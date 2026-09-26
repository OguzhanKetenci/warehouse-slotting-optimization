"""
Module 3 - KPI Summary Workbook

Input : data/processed/scenario_results.csv, data/processed/slot_assignment_velocity.csv,
        data/processed/clean_lines.csv   (all produced by Modules 1-2)
Run   : python src/03_kpi_summary.py   (from repo root, after Modules 1 and 2)
Output: reports/kpi_summary.xlsx
  Sheet "KPI Summary"  - one row per scenario
  Sheet "Assumptions"  - every parameter and modelling assumption
  Sheet "SKU Slotting" - slot assignment of the 100 fastest SKUs (full-velocity scenario)

The workbook numbers are read from scenario_results.csv without rounding and checked
against the CSV after saving.
"""
import importlib.util
import math
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
PROC = ROOT / "data" / "processed"
OUT = ROOT / "reports" / "kpi_summary.xlsx"

WALK_SPEED_M_PER_S = 1.0     # assumed average picker walking speed
TOP_N_SKUS = 100

HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
HEADER_FONT = Font(bold=True)
THIN = Side(style="thin", color="9AA5B1")


def load_module2():
    """Import Module 2 so the Assumptions sheet reads the same constants the model used."""
    spec = importlib.util.spec_from_file_location("slotting_routing", ROOT / "src" / "02_slotting_routing.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


KPI_COLUMNS = [
    # (header, csv column or None, number format, width)
    ("Scenario", "scenario", "@", 22),
    ("Runs (seeds)", "n_runs", "0", 13),
    ("Avg. distance per order (m)", "avg_distance_per_order_m", "#,##0.0", 18),
    ("Min across seeds (m)", "avg_distance_per_order_min_m", "#,##0.0", 16),
    ("Max across seeds (m)", "avg_distance_per_order_max_m", "#,##0.0", 16),
    ("Total annual distance (km)", "total_distance_km", "#,##0.0", 18),
    ("Change vs. random baseline", "change_vs_random", "0.0%", 18),
    ("Est. annual walking time (h)", None, "#,##0", 18),
    ("Picks from nearest 20% of slots", "share_picks_nearest_20pct_slots", "0.0%", 20),
]


def style_header(ws, row: int, n_cols: int) -> None:
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = Border(bottom=THIN)


def build_kpi_sheet(ws, results: pd.DataFrame, n_orders: int, n_runs_random: int) -> None:
    ws.title = "KPI Summary"
    for c, (header, _, _, width) in enumerate(KPI_COLUMNS, start=1):
        ws.cell(row=1, column=c, value=header)
        ws.column_dimensions[get_column_letter(c)].width = width
    style_header(ws, 1, len(KPI_COLUMNS))
    ws.row_dimensions[1].height = 34
    walk_col = [h for h, *_ in KPI_COLUMNS].index("Est. annual walking time (h)") + 1
    ws.cell(row=1, column=walk_col).comment = Comment(
        f"Assumption: average walking speed {WALK_SPEED_M_PER_S} m/s. "
        "Walking time = total distance / speed. Excludes picking, handling, searching and congestion.",
        "Module 3", width=320, height=90)

    for r, rec in enumerate(results.to_dict("records"), start=2):
        for c, (_, key, fmt, _) in enumerate(KPI_COLUMNS, start=1):
            value = rec[key] if key else rec["total_distance_km"] * 1000.0 / WALK_SPEED_M_PER_S / 3600.0
            if key == "n_runs":
                value = int(value)
            cell = ws.cell(row=r, column=c, value=value)
            cell.number_format = fmt
            if c > 1:
                cell.alignment = Alignment(horizontal="right")
    ws.freeze_panes = "B2"

    notes = [
        f"Random baseline and Class-based ABC are means over {n_runs_random} random seeds; min/max columns show the range across seeds. "
        "Full velocity is deterministic (1 run).",
        f"Total distance = sum of S-shape route lengths over all {n_orders:,} orders of the last 12 months, "
        "on a synthetic single-block warehouse (see Assumptions).",
        f"Estimated walking time = total distance / {WALK_SPEED_M_PER_S} m/s. It is an indicator, not a labour-hours forecast.",
    ]
    for i, text in enumerate(notes):
        cell = ws.cell(row=len(results) + 3 + i, column=1, value=text)
        cell.font = Font(italic=True, color="52514E")


def build_assumptions_sheet(ws, sr, n_orders: int, n_skus: int, n_visits: int, near_cut: int) -> None:
    ws.title = "Assumptions"
    n_slots = sr.N_AISLES * sr.SIDES_PER_AISLE * sr.SLOTS_PER_SIDE
    rows = [
        ("Data", None, None),
        ("Source", "UCI Online Retail II, sheet 'Year 2010-2011'", "Real order data; warehouse layout is synthetic."),
        ("Period evaluated", "Last 12 months, all orders", "No sampling; every order is routed."),
        ("Orders evaluated", n_orders, "One invoice = one picking task."),
        ("SKUs slotted", n_skus, "Valid product codes only (cancellations, fees and non-product codes removed)."),
        ("Order-SKU visits", n_visits, "Distinct (order, SKU) pairs; two lines of one SKU in an order = one visit to the slot."),
        ("Velocity metric", "Pick frequency (order lines per SKU)", "From Module 1; not revenue."),
        ("ABC thresholds", "A: 80% of picks, B: next 15%, C: last 5%", "Cumulative pick-frequency cut-offs from Module 1."),
        ("Layout", None, None),
        ("Aisles", sr.N_AISLES, "Parallel aisles, single block."),
        ("Racks per aisle", sr.SIDES_PER_AISLE, "Left and right side."),
        ("Slots per rack", sr.SLOTS_PER_SIDE, None),
        ("Total slots", n_slots, f"{n_slots - n_skus:,} spare slots."),
        ("Slot width (m)", sr.SLOT_WIDTH_M, None),
        ("Aisle length (m)", sr.AISLE_LENGTH_M, "Front cross aisle to back cross aisle."),
        ("Aisle pitch (m)", sr.AISLE_PITCH_M, "Distance between aisle centre lines."),
        ("First aisle centre x (m)", sr.FIRST_AISLE_X_M, None),
        ("Depot", "Left end of front cross aisle (x = 0, y = 0)", "Pick-up and drop-off point of every route."),
        ("Distance metric", "Rectilinear, metres", "Cross-aisle width and lateral movement inside an aisle are ignored."),
        ("Routing", None, None),
        ("Heuristic", "S-shape", "Every aisle with a pick is traversed end to end; if the number of such aisles is odd, "
                                 "the last one is entered and left from its farthest pick; route returns to the depot."),
        ("Scenarios", None, None),
        ("Random (baseline)", f"{len(sr.RANDOM_SEEDS)} seeds", "SKUs assigned to a random subset of slots. Baseline = mean of the seeds."),
        ("Class-based ABC", f"{len(sr.RANDOM_SEEDS)} seeds", "A SKUs in the nearest slots, then B, then C; random within each zone."),
        ("Full velocity", "Deterministic", "Fastest SKU in the nearest slot; slots ranked by walking distance from the depot."),
        ("Random seeds", ", ".join(str(s) for s in sr.RANDOM_SEEDS), None),
        ("KPI definitions", None, None),
        ("Near slots", f"Nearest {sr.NEAR_SLOT_SHARE:.0%} of slots by walking distance ({near_cut:,} slots)",
         "Share of picks served from these slots."),
        ("Walking speed (m/s)", WALK_SPEED_M_PER_S, "Used only for the estimated walking time."),
        ("Annual distance", "Sum over the 12-month data window", "No growth or seasonality adjustment."),
        ("Not modelled", None, None),
        ("Out of scope", "Item size/weight, rack capacity, congestion, batching, replenishment, re-slotting cost",
         "Results are a first-order estimate of travel savings, not a full operational forecast."),
    ]
    ws.append(["Parameter", "Value", "Note"])
    style_header(ws, 1, 3)
    for name, value, note in rows:
        ws.append([name, value, note])
        row = ws.max_row
        if value is None and note is None:          # section header
            ws.cell(row=row, column=1).font = Font(bold=True)
            for c in (1, 2, 3):
                ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor="F0EFEC")
        else:
            ws.cell(row=row, column=2).alignment = Alignment(horizontal="left", wrap_text=True, vertical="top")
            ws.cell(row=row, column=3).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=row, column=1).alignment = Alignment(vertical="top")
            if isinstance(value, (int, float)):
                ws.cell(row=row, column=2).number_format = "#,##0.0" if isinstance(value, float) else "#,##0"
    for col, width in zip("ABC", (26, 46, 80)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"


SKU_COLUMNS = [
    ("Velocity rank", "velocity_rank", "0", 14), ("StockCode", "StockCode", "@", 12),
    ("Description", "description", "@", 38), ("Pick lines", "pick_lines", "#,##0", 12),
    ("ABC class", "abc_pick", "@", 10), ("Aisle", "aisle", "0", 8), ("Side", "side", "@", 8),
    ("Position", "position", "0", 10), ("Walking distance from depot (m)", "walk_distance_m", "#,##0.0", 18),
]


def build_sku_sheet(ws, top: pd.DataFrame) -> None:
    ws.title = "SKU Slotting"
    for c, (header, _, _, width) in enumerate(SKU_COLUMNS, start=1):
        ws.cell(row=1, column=c, value=header)
        ws.column_dimensions[get_column_letter(c)].width = width
    style_header(ws, 1, len(SKU_COLUMNS))
    ws.row_dimensions[1].height = 34
    for r, rec in enumerate(top.to_dict("records"), start=2):
        for c, (_, key, fmt, _) in enumerate(SKU_COLUMNS, start=1):
            value = rec[key]
            if hasattr(value, "item"):
                value = value.item()                 # numpy scalar -> python scalar
            cell = ws.cell(row=r, column=c, value=value)
            cell.number_format = fmt
    ws.freeze_panes = "A2"


def verify_workbook(path: Path, results: pd.DataFrame, top: pd.DataFrame) -> None:
    """Re-open the saved workbook and require exact equality with the CSV values."""
    wb = load_workbook(path)
    ws = wb["KPI Summary"]
    for r, rec in enumerate(results.to_dict("records"), start=2):
        for c, (header, key, _, _) in enumerate(KPI_COLUMNS, start=1):
            cell = ws.cell(row=r, column=c).value
            if key is None:   # derived column: openpyxl stores 16 significant digits
                expected = rec["total_distance_km"] * 1000.0 / WALK_SPEED_M_PER_S / 3600.0
                ok = math.isclose(cell, expected, rel_tol=1e-12)
            else:             # CSV columns must match exactly
                expected = int(rec[key]) if key == "n_runs" else rec[key]
                ok = cell == expected
            assert ok, f"KPI Summary row {r} '{header}': {cell!r} != {expected!r}"
    ws = wb["SKU Slotting"]
    for r, rec in enumerate(top.to_dict("records"), start=2):
        for c, (header, key, _, _) in enumerate(SKU_COLUMNS, start=1):
            cell, expected = ws.cell(row=r, column=c).value, rec[key]
            assert cell == expected, f"SKU Slotting row {r} '{header}': {cell!r} != {expected!r}"
    print(f"Verification OK: {len(results)} scenario rows x {len(KPI_COLUMNS)} columns and "
          f"{len(top)} SKU rows match the CSV files exactly.")


def main() -> None:
    sr = load_module2()
    results = pd.read_csv(PROC / "scenario_results.csv")
    assignment = pd.read_csv(PROC / "slot_assignment_velocity.csv", dtype={"StockCode": str})
    top = assignment.sort_values("velocity_rank").head(TOP_N_SKUS).reset_index(drop=True)

    lines = pd.read_csv(PROC / "clean_lines.csv", usecols=["Invoice", "StockCode"], dtype=str)
    n_orders, n_visits = lines["Invoice"].nunique(), len(lines.drop_duplicates())
    n_slots = sr.N_AISLES * sr.SIDES_PER_AISLE * sr.SLOTS_PER_SIDE
    near_cut = int(round(sr.NEAR_SLOT_SHARE * n_slots))

    wb = Workbook()
    build_kpi_sheet(wb.active, results, n_orders, int(results["n_runs"].max()))
    build_assumptions_sheet(wb.create_sheet(), sr, n_orders, len(assignment), n_visits, near_cut)
    build_sku_sheet(wb.create_sheet(), top)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    verify_workbook(OUT, results, top)
    print(f"Workbook -> {OUT}")
    view = results.assign(est_walking_hours=results["total_distance_km"] * 1000 / WALK_SPEED_M_PER_S / 3600)
    print(view[["scenario", "avg_distance_per_order_m", "total_distance_km", "change_vs_random",
                "est_walking_hours", "share_picks_nearest_20pct_slots"]].to_string(index=False, float_format=lambda v: f"{v:,.4f}"))


if __name__ == "__main__":
    main()
